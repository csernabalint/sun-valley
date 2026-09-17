import docx

def dump_docx(filename, out_filename):
    doc = docx.Document(filename)
    with open(out_filename, 'w', encoding='utf-8') as f:
        f.write(f"=== {filename} ===\n\n")
        for p in doc.paragraphs:
            if p.text.strip():
                f.write(p.text + "\n")
        f.write("\n--- TABLES ---\n")
        for t in doc.tables:
            for row in table_rows(t):
                f.write(" | ".join(row) + "\n")
            f.write("\n")

def table_rows(table):
    rows = []
    for r in table.rows:
        rows.append([c.text.replace('\n', ' ').strip() for c in r.cells])
    return rows

dump_docx('docs/Sun Valley Weboldal Specifikáció.docx', 'docs/spec_dump.txt')
dump_docx('docs/Sun Valley Zrt. Céginformációk.docx', 'docs/ceg_dump.txt')
print("Successfully dumped docx files to text.")
